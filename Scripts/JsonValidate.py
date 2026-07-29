import json
import sqlite3
import re
import os
import pdfplumber
import pandas as pd
from jsonschema import Draft7Validator, FormatChecker
from MainModule import JsonOperations,APIOperations
from datetime import datetime,date
from openpyxl.styles import Alignment, PatternFill
from openpyxl import load_workbook
from pathlib import Path


class C3_MPP_JsonSchema():
    def __init__(self,ProjectsPath,Product,Mode):
        
        self.Projects=ProjectsPath
        self.Product = Product
        self.Mode = Mode
        self.Header = {}
        self.Header['Product'] = self.Product
        self.Header['Mode'] = self.Mode
         # Connect to the SQLite database
        self.Conn =  sqlite3.connect('Resources/TestDataFile.db', check_same_thread=False)
        self.Japi = JsonOperations('json/Xpath.json')
        JapiDatatemp =self.Japi.read_file()
        self.JapiData = JapiDatatemp['API']


    def SegregateProjectFolder(self):
        self.TClist={}
        for pro in self.Projects:           
            Backuppath = None
            jsonpath = None
            for root, dirs, files in os.walk(pro):
                if "ReferenceData" in dirs:
                    dirs.remove("ReferenceData")
                for file in files:
                    if file.endswith("Final_TestBackup.gproj"):
                        Backuppath =os.path.join(Path(root), Path(file))
                    if file.endswith("FinalReport.json"):
                        jsonpath = os.path.join(Path(root), Path(file))
            #Fetch Testcases from Backup json file
            BKJSON = JsonOperations(Backuppath)
            self.BKJSONData =  BKJSON.read_file()
            ReportJSON = JsonOperations(jsonpath)
            self.ReportJsonData =  ReportJSON.read_file()
            self.ExtractTestFromDB()
            for TCdata in self.BKJSONData['testBkpTestResultsandPath']:
                tracepathlist = TCdata['actualTracePath'].split('\\')
                ExactPath=Backuppath.replace('\\', '/').rsplit('/', 1)[0]# Remove the last segment
                tracepath = os.path.join(ExactPath,tracepathlist[len(tracepathlist)-3],tracepathlist[len(tracepathlist)-2],tracepathlist[len(tracepathlist)-1])
                TestId,TestName,TestResult=TCdata['testinformation']['TestId'],TCdata['testinformation']['TestName'],TCdata['testinformation']['TestResult']

                TraceQuery= f'''
                                UPDATE JsonDetails
                                SET JTracepath='{str(tracepath)}', TestResult='{TestResult}'
                                WHERE JTestID='{TestId}'AND JTestName='{TestName}'
                            '''

                self.ExecutebyQuery(TraceQuery)
            self.ExtractFromTrace()



    def ExtractTestFromDB(self):

        self.Certification=self.BKJSONData['testBkpProjectConfiguration']['EsdfConfigurationModel']['AllESDFFields']['SpecificationSupported']
        Query= f"SELECT CTSID, HWID FROM CTSVERSIONS WHERE Product='{self.Product}' AND Mode='{self.Mode}' AND Specification ='{self.Certification}'"
        res=self.FetchDataFromQRY(Query)
        CTS_ID,HW_ID=res['CTSID'][0],res['HWID'][0]
      
        FetchQuery=  f'''
                    SELECT 
                        TestCaseDetails.TEST_DESCRIPTION,
                        TestCaseCompleteDetails.TESTCASE_ID,
                        CoilType.COIL_TYPE,
                        CoilType.COIL_VOLTAGE,
                        CoilType.COIL_MODULATION,
                        CoilType.COIL_LOAD,
                        CoilType.PT_PHASE_LOAD
                    FROM TestCaseCompleteDetails
                    LEFT JOIN TestCaseDetails ON TestCaseCompleteDetails.TEST_NAME = TestCaseDetails.ID
                    LEFT JOIN CoilType ON TestCaseCompleteDetails.COIL_TYPE = CoilType.ID
                    WHERE 
                        TestCaseCompleteDetails.HW_TYPE = "{HW_ID}" 
                        AND TestCaseCompleteDetails.CTS_VERSION = "{CTS_ID}"
                '''
        res=self.FetchDataFromQRY(FetchQuery)

        self.ExecutebyQuery(f"DELETE FROM JsonDetails")

        if res is not None:
            for _, row in res.iterrows():
                insert_query = f"""
                    INSERT OR REPLACE INTO JsonDetails (JTestID, JTestName, JCoilType,JCoilVoltage,JCoilModulation,JCoilLoad,JPtPhaseLoad)
                    VALUES ('{row['TESTCASE_ID']}', '{row['TEST_DESCRIPTION']}', '{row['COIL_TYPE']}','{row['COIL_VOLTAGE']}','{row['COIL_MODULATION']}','{row['COIL_LOAD']}','{row['PT_PHASE_LOAD']}')
                """
                self.ExecutebyQuery(insert_query)

      
    def ExtractFromTrace(self):

        Results={}
        TestCount=0
        for Test in self.ReportJsonData['TestExecutionDetails']['TestScope']: 
            if Test not in Results:Results[Test]={}

            GetTestDetailsQuery=f'''   SELECT * FROM JsonDetails WHERE JTestName='{Test}' '''
            res=self.FetchDataFromQRY(GetTestDetailsQuery)
            
            if res is not None:

                self.TraceUPL = APIOperations(url=self.JapiData[self.Product][self.Mode]['PutWaveformFile'])    
                self.TraceUPL.files = {"WaveformFile":open(res['JTracePath'][0].replace('/','\\'),"rb")}
                status = self.TraceUPL.PutRequest()
                if status == 200:
                    print(f'trace is loaded')
                     #Get Packets___________________________________________________________________________
                    self.PktAPI = APIOperations(url=self.JapiData[self.Product][self.Mode]['GetCCLinePackets'],retype='json')
                    self.file_list = self.PktAPI.GetRequest()
                    id=0
                    PktCount=0
                    if 'TestingScope' not in Results[Test]:Results[Test]['TestingScope']={}
                    if 'TestLogs' not in Results[Test]['TestingScope']:Results[Test]['TestingScope']['TestLogs']={}
                    if 'TesterConfiguration' not in Results[Test]['TestingScope']:Results[Test]['TestingScope']['TesterConfiguration']={}

                    
                    try:
                        while id < len(self.file_list):
                            Type=self.GetPacketType(id)
                            if Type in ['Packet','Response']:
                                raw_data = [int(x, 16) for x in self.file_list[id]['header_Payload']['sFieldType'].split(':')[1].split()]
                                Pktduration=round((self.file_list[id]['stopTime']-self.file_list[id]['startTime'])*1000,1)
                                PktDescription = self.PktDescription(id)  
                                # print(PktDescription) 
                                if id not in Results[Test]['TestingScope']['TestLogs']:Results[Test]['TestingScope']['TestLogs'][id]={}
                                if self.ReportJsonData['TestingScope'][TestCount]['TestLogs'][PktCount]['PacketDescription']!=PktDescription:
                                    Results[Test]['TestingScope']['TestLogs'][id]['PacketDescription']=f'Received PktDesc from Trace is {PktDescription} and Json Description is { self.ReportJsonData['TestingScope'][TestCount]['TestLogs'][PktCount]['PacketDescription']}'
                                if round(self.ReportJsonData['TestingScope'][TestCount]['TestLogs'][PktCount]['PacketDuration'],1) !=Pktduration:
                                    Results[Test]['TestingScope']['TestLogs'][id]['PacketDuration']=f'Received PktDuration from Trace is {Pktduration} and Json Duration is {round(self.ReportJsonData['TestingScope'][TestCount]['TestLogs'][PktCount]['PacketDuration'],1)}'
                                if self.ReportJsonData['TestingScope'][TestCount]['TestLogs'][PktCount]['RawData']!=raw_data:
                                    Results[Test]['TestingScope']['TestLogs'][id]['RawData']=f'Received RawData from Trace is {raw_data} and Json RawData is { self.ReportJsonData['TestingScope'][TestCount]['TestLogs'][PktCount]['RawData']}'
                                PktCount+=1
                                if Results[Test]['TestingScope']['TestLogs'][id]=={}:del Results[Test]['TestingScope']['TestLogs'][id]
                            id+=1
                    except Exception as e: print(e)
            else:Results[Test]="Test did not found in DataBase"
            TestCount+=1
            
        
        with open('JsonResults.json', 'w') as json_file:
            json.dump(Results, json_file, indent=4)
    def PktDescription(self,id):

        pkt= self.file_list[id]['pktType']

        if 'SADT' in pkt:
            match = re.match(r'SADT/(\d+)([eo])', pkt)
            if match:
                number = match.group(1)
                parity = "even" if match.group(2) == 'e' else "odd"
                return f"Simultaneous_Auxiliary_Data_Transport_{number}_{parity} "

        elif pkt not in ['SADC','SDSR','DSR',"SRQ [0x20] "]:
            if pkt in ['Signal strength']: return f'{pkt.replace(' ','_').replace(':','_').title()} {self.file_list[id]['value']}'
            return f'{pkt.replace(' ','_').replace(':','_')} {self.file_list[id]['value']}'
        else:
            if pkt=='SADC':
                pkt='Simultaneous_Auxilory_Data_Control'
            elif pkt=='SDSR':
                pkt='Simultaneous_Data_Stream_Response'
            elif pkt=='DSR':
                pkt='Data_Stream_Response'
            elif pkt=="SRQ [0x20] ":
                pkt='Specific_Request'
        return f'{pkt} {self.file_list[id]['value']}'                        


    #1.Get packet Type, testermsg/packet/response
    def GetPacketType(self,id):
        if self.Header['Product'] == "C3":
            if self.Header['Mode'] == 'TPT':
                if self.file_list[id]['isTesterPkt']==False and self.file_list[id]['isFWTestermessage']==False:
                    return 'Packet'
                elif self.file_list[id]['isTesterPkt']==True and self.file_list[id]['isFWTestermessage']==False:
                    return 'Response'
                elif self.file_list[id]['isTesterPkt']==True and self.file_list[id]['isFWTestermessage']==True:
                    return 'TesterMsg'
            elif self.Header['Mode'] =="TPR":
                if self.file_list[id]['isTesterPkt']==True and self.file_list[id]['isFWTestermessage']==False:
                    return 'Packet'
                elif self.file_list[id]['isTesterPkt']==False and self.file_list[id]['isFWTestermessage']==False:
                    return 'Response'
                elif self.file_list[id]['isTesterPkt']==True and self.file_list[id]['isFWTestermessage']==True:
                    return 'TesterMsg'
        elif self.Header['Product'] == "MPP":
            if self.Header['Mode']=="TPR":
                if self.file_list[id]['isTesterPkt']==True and self.file_list[id]['isFWTestermessage']==False:
                    return 'Packet'
                elif self.file_list[id]['isTesterPkt']==False and self.file_list[id]['isFWTestermessage']==False:
                    return 'Response'
                elif self.file_list[id]['isTesterPkt']==True and self.file_list[id]['isFWTestermessage']==True:
                    return 'TesterMsg'
            elif self.Header['Mode']=="TPT":
                if self.file_list[id]['isTesterPkt']==True and self.file_list[id]['isFWTestermessage']==False:
                    return 'Packet'
                elif self.file_list[id]['isTesterPkt']==False and self.file_list[id]['isFWTestermessage']==False:
                    return 'Response'
                elif self.file_list[id]['isTesterPkt']==True and self.file_list[id]['isFWTestermessage']==True:
                    return 'TesterMsg'
        return None
            
    def FetchDataFromQRY(self,QRY):
        try:
            QRY_DF = pd.read_sql_query(QRY, self.Conn)
            if QRY_DF.shape[0] > 0:
                # print(list(QRY_DF['Name']))
                return QRY_DF
            return None
        except Exception as e:
            print(e)
            return None
    def ExecutebyQuery(self, QRY, params=None):
        """
        Execute a SQL query with optional parameters.

        :param QRY:   SQL query string (may contain ? placeholders)
        :param params: tuple or list of parameters for placeholders
        """
        try:
            self.cursor = self.Conn.cursor()
            if params:
                self.cursor.execute(QRY, params)
            else:
                self.cursor.execute(QRY)
            self.Conn.commit()
        except sqlite3.Error as e:
            print(f"Error occurred: {e}")
        except Exception as e:
            print(e)


        

obj=C3_MPP_JsonSchema(
    ProjectsPath=[r"C:\Users\Eswar\Downloads\Apple_V221_120226_173255 1"],
    Product='MPP',Mode='TPR'
)
obj.SegregateProjectFolder()